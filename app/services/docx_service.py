import io
import os
import docx
from docx import Document
from docx.shared import Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from typing import Dict, Any

class DocxService:
    @staticmethod
    def _get_document() -> Document:
        try:
            return Document()
        except Exception:
            templates_dir = os.path.join(os.path.dirname(docx.__file__), 'templates')
            default_template = os.path.join(templates_dir, 'default-docx-template')
            if os.path.exists(default_template):
                return Document(default_template)
            raise

    @staticmethod
    def generate_docx(data: Dict[str, Any]) -> io.BytesIO:
        doc = DocxService._get_document()

        style = doc.styles['Normal']
        font = style.font
        font.name = 'Arial'
        font.size = Pt(11)

        # Title
        title_p = doc.add_paragraph()
        title_run = title_p.add_run("CLAUDE AI CHANNEL CLONE — AUTOMATED BRIEF & SCRIPT")
        title_run.bold = True
        title_run.font.size = Pt(20)
        title_run.font.color.rgb = RGBColor(139, 92, 246)
        title_p.alignment = WD_ALIGN_PARAGRAPH.CENTER

        doc.add_paragraph("----------------------------------------------------------------------------------")

        # STAGE 2: BRANDING BRIEF
        doc.add_heading("STATE 2 — BRANDING BRIEF", level=1)
        brief = data.get("stage2_branding_brief") or {}

        doc.add_heading("Suggested Channel Name Variants:", level=2)
        for name in (brief.get("suggested_channel_names") or []):
            doc.add_paragraph(f"• {str(name)}")

        doc.add_heading("Channel Description Variants:", level=2)
        for desc in (brief.get("channel_description_variants") or []):
            p = doc.add_paragraph()
            p.add_run(f"\"{str(desc)}\"").italic = True

        doc.add_heading("Branding Image Prompts:", level=2)
        p1 = doc.add_paragraph()
        p1.add_run("Logo Prompt: ").bold = True
        p1.add_run(str(brief.get("logo_generation_prompt") or "N/A"))

        p2 = doc.add_paragraph()
        p2.add_run("Banner Prompt: ").bold = True
        p2.add_run(str(brief.get("banner_generation_prompt") or "N/A"))

        doc.add_page_break()

        # STAGE 5: STYLE DNA
        doc.add_heading("STATE 5 — ANALYSIS & STYLE DNA", level=1)
        dna = data.get("stage5_style_dna") or data.get("original_analysis") or {}
        if isinstance(dna, dict):
            for k, v in dna.items():
                p = doc.add_paragraph()
                p.add_run(f"{str(k).replace('_', ' ').title()}: ").bold = True
                p.add_run(str(v or "N/A"))

        doc.add_page_break()

        # STAGE 6: FULL SCRIPT
        doc.add_heading("STATE 6 — FULL SCRIPT GENERATION", level=1)
        script_data = data.get("stage6_script") or data.get("cloned_content") or {}
        if isinstance(script_data, dict):
            doc.add_paragraph(f"Target Word Count: {str(script_data.get('target_word_count') or 'N/A')} | Final Word Count: {str(script_data.get('final_word_count') or 'N/A')}")

            doc.add_heading("3-Second Opening Hook:", level=2)
            hp = doc.add_paragraph()
            hp.add_run(str(script_data.get("hook") or "N/A")).bold = True

            doc.add_heading("Full Teleprompter Script:", level=2)
            sp = doc.add_paragraph()
            sp.add_run(str(script_data.get("full_script") or script_data.get("script") or "N/A"))

            doc.add_heading("Outro / Call to Action:", level=2)
            doc.add_paragraph(str(script_data.get("outro") or script_data.get("call_to_action") or "N/A"))

        doc.add_page_break()

        # STAGE 8 & 9: BEAT-BY-BEAT VISUAL & VIDEO PROMPTS
        doc.add_heading("STATE 8 & 9 — SCRIPT BEAT PROMPTS (3-5s each)", level=1)
        beats = data.get("stage7_8_9_beat_prompts") or data.get("scene_visuals") or []

        if isinstance(beats, list):
            for beat in beats:
                if isinstance(beat, dict):
                    b_p = doc.add_paragraph()
                    b_p.add_run(f"Beat #{str(beat.get('beat_num') or 1)} | Segment: ").bold = True
                    b_p.add_run(f"\"{str(beat.get('script_segment') or beat.get('section_name') or '')}\"").italic = True

                    doc.add_paragraph(f"• Image Prompt: {str(beat.get('image_prompt') or beat.get('ai_image_prompt') or 'N/A')}")
                    doc.add_paragraph(f"• Video Motion Prompt: {str(beat.get('video_motion_prompt') or 'N/A')}")
                    doc.add_paragraph(f"• Camera: {str(beat.get('camera_angle') or 'Wide')} | Lighting: {str(beat.get('lighting') or 'Cinematic')} | Mood: {str(beat.get('mood') or 'Engaging')}")
                    doc.add_paragraph("---")

        doc.add_page_break()

        # STAGE 11: THUMBNAILS
        doc.add_heading("STATE 11 — THUMBNAIL CONCEPTS", level=1)
        thumbs = data.get("stage10_11_thumbnails") or ([data.get("thumbnail_concept")] if data.get("thumbnail_concept") else [])
        if isinstance(thumbs, list):
            for idx, t in enumerate(thumbs, 1):
                if isinstance(t, dict):
                    doc.add_heading(f"Thumbnail Concept #{idx}: {str(t.get('title') or 'Concept')}", level=2)
                    doc.add_paragraph(f"• Visual Concept: {str(t.get('visual_concept') or t.get('visual_description') or 'N/A')}")
                    doc.add_paragraph(f"• Text Overlay: {str(t.get('text_overlay') or 'N/A')}")
                    doc.add_paragraph(f"• Emotion Trigger: {str(t.get('emotion_trigger') or 'N/A')}")
                    doc.add_paragraph(f"• Midjourney Prompt: {str(t.get('midjourney_prompt') or t.get('ai_image_prompt') or 'N/A')}")

        target_stream = io.BytesIO()
        doc.save(target_stream)
        target_stream.seek(0)
        return target_stream
